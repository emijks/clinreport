import io
import shutil
import sqlite3
from pathlib import Path

import pytest
from docx import Document


def docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return '\n'.join(parts)


def test_open_lists_samples(service, trio_db, single_db):
    trio = service.open_result({'dbpath': trio_db})
    assert [s['label'] for s in trio['samples']] == ['sample1', 'sample2', 'sample3']
    assert trio['settings']['db_configured'] is False

    # a single-sample result without sample ids is labelled by the file name
    single = service.open_result({'dbpath': single_db})
    assert [s['label'] for s in single['samples']] == ['pah2']


def test_open_rejects_missing_and_foreign_files(service, tmp_path):
    with pytest.raises(service.ClientError):
        service.open_result({})
    with pytest.raises(service.ClientError):
        service.open_result({'dbpath': str(tmp_path / 'nope.sqlite')})
    foreign = tmp_path / 'foreign.sqlite'
    sqlite3.connect(foreign).execute('create table t (x)').connection.close()
    with pytest.raises(service.ClientError):
        service.open_result({'dbpath': str(foreign)})


def test_resolve_by_job_id(service, trio_db, tmp_path):
    job_dir = Path(service._admin_util().get_jobs_dir()) / '260310-121214'
    job_dir.mkdir(parents=True)
    shutil.copy(trio_db, job_dir / 'trio1.sqlite')
    assert service.resolve_dbpath(job_id='260310-121214') == str(job_dir / 'trio1.sqlite')


def test_resolve_by_job_id_in_user_folder_without_username(service, trio_db):
    # OpenCRAVAT 3.x keeps even single-user jobs in <jobs>/default/<id>, and its Jobs page
    # opens the viewer with only ?job_id=, so the widget cannot pass a username.
    job_dir = Path(service._admin_util().get_jobs_dir()) / 'default' / '260925-001732'
    job_dir.mkdir(parents=True)
    shutil.copy(trio_db, job_dir / 'trio1.sqlite')
    assert service.resolve_dbpath(job_id='260925-001732') == str(job_dir / 'trio1.sqlite')
    assert service.resolve_dbpath(job_id='260925-001732', username='default') == str(job_dir / 'trio1.sqlite')


@pytest.mark.parametrize('job_id, username', [('..', ''), ('a/b', ''), ('a\\b', ''), ('x', '..'), ('x', 'u/v')])
def test_resolve_rejects_path_like_job_ids(service, job_id, username):
    with pytest.raises(service.ClientError, match='Некорректный'):
        service.resolve_dbpath(job_id=job_id, username=username)


def test_resolve_unknown_job(service):
    with pytest.raises(service.ClientError, match='не найдено'):
        service.resolve_dbpath(job_id='nope')


def test_process_default_builds_a_report_per_sample(service, trio_db):
    result = service.process({'dbpath': trio_db, 'target': 0, 'mode': 'default'})
    reports = result['reports']
    assert [(r['label'], r['kind'], r['is_target']) for r in reports] == [
        ('sample1', 'default', True),
        ('sample2', 'default', False),
        ('sample3', 'default', False),
    ]
    first = reports[0]['context']
    assert len(first['p_variants']) == 1 and len(first['lp_variants']) == 1 and len(first['vus_variants']) == 1
    assert reports[0]['filename'] == 'Заключение (sample1).docx'
    assert set(result['layouts']) == {'default', '10x'}


def test_process_lpwgs_target_default_others_10x(service, trio_db):
    reports = service.process({'dbpath': trio_db, 'target': 1, 'mode': 'lpwgs'})['reports']
    assert [(r['label'], r['kind']) for r in reports] == [
        ('sample2', 'default'),
        ('sample1', '10x'),
        ('sample3', '10x'),
    ]
    assert reports[1]['filename'] == 'Заключение LPWGS (sample1).docx'
    # the 10x table lists the target's causative variants (P + LP here)
    assert len(reports[1]['context']['main_variants']) == 2


def test_process_uses_clinician_from_settings(service, single_db):
    service.update_settings({'clinician': 'Иванова И.И.'})
    context = service.process({'dbpath': single_db})['reports'][0]['context']
    assert context['clinician'] == 'Иванова И.И.'


def test_sample_without_id_is_numbered_by_file_name(service, single_db):
    context = service.process({'dbpath': single_db})['reports'][0]['context']
    assert context['patient']['id'] == 'pah2'  # not the literal 'None'


@pytest.mark.parametrize('template', ['DZM', 'FND'])
def test_render_default_templates_with_edits(service, single_db, template):
    report = service.process({'dbpath': single_db})['reports'][0]
    context = report['context']
    context['patient']['diagnosis'] = 'Фенилкетонурия (проверка правки)'
    context['p_variants'][0]['disease'] = 'Отредактированное заболевание'

    result = service.render(
        {'context': context, 'kind': 'default', 'template': template, 'filename': report['filename']}
    )
    assert result.content_type == service.DOCX_MIME
    assert result.filename == 'Заключение (pah2).docx'
    text = docx_text(result.data)
    assert 'Фенилкетонурия (проверка правки)' in text
    assert 'Отредактированное заболевание' in text
    assert context['p_variants'][0]['gene'] in text


def test_render_10x(service, trio_db):
    report = service.process({'dbpath': trio_db, 'mode': 'lpwgs'})['reports'][1]
    result = service.render({'context': report['context'], 'kind': '10x', 'filename': report['filename']})
    text = docx_text(result.data)
    assert report['context']['main_variants'][0]['hgvsg'] in text


def test_render_rejects_unknown_template(service, single_db):
    context = service.process({'dbpath': single_db})['reports'][0]['context']
    with pytest.raises(service.ClientError):
        service.render({'context': context, 'template': '../../evil'})


def test_settings_roundtrip(service):
    saved = service.update_settings({'clinician': ' Иванова ', 'auto_upload': True, 'unknown': 'x'})
    assert saved['clinician'] == 'Иванова' and saved['auto_upload'] is True
    assert 'unknown' not in service.get_settings({})


def test_upload_and_similar_need_configured_db(service, single_db):
    context = service.process({'dbpath': single_db})['reports'][0]['context']
    with pytest.raises(service.ClientError):
        service.upload({'context': context})
