"""Custom report templates, the template check, ru_status/ru_refresh and the desktop edition."""

import asyncio
import base64
import io
import os
import zipfile

import pytest
from clinreport_core import templates as tpl
from docx import Document


def docx_with(text: str, base='DZM') -> bytes:
    """A built-in template with one more paragraph containing `text`."""
    document = Document(str(tpl.BUILTIN_DIR / f'{base}.docx'))
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def add(service, title, text='x', kind='default', base='DZM', caller=None, **extra):
    params = {'title': title, 'kind': kind, 'data': b64(docx_with(text, base)), **extra}
    return service.add_template(params, caller or service.LOCAL)['template']


@pytest.mark.parametrize(('template_id', 'kind'), [(tid, kind) for tid, _title, kind in tpl.BUILTIN])
def test_builtin_templates_use_only_known_fields(template_id, kind):
    assert tpl.check(tpl.BUILTIN_DIR / f'{template_id}.docx', kind) == {'ok': True, 'error': None, 'unknown': []}


def test_check_names_misspelled_fields_with_suggestions():
    data = docx_with(
        '{{ patient.birth_dat }} {{ patient.diagnosis }} {% for v in p_variants %}{{ v.gen }}{% endfor %} {{ clinicin }}'
    )
    result = tpl.check(data, 'default')
    assert not result['ok'] and result['error'] is None
    found = {u['field']: u['suggestion'] for u in result['unknown']}
    assert found['p_variants[].gen'] == 'p_variants[].gene'
    assert found['clinicin'] == 'clinician'
    assert 'patient.birth_dat' in found and 'patient.diagnosis' not in found


def test_check_reports_syntax_errors_and_non_docx():
    assert 'Ошибка в полях шаблона' in tpl.check(docx_with('{% for v in p_variants %}'), 'default')['error']
    assert '.docx' in tpl.check(b'not a zip', 'default')['error']


def _paths(value, prefix=''):
    if isinstance(value, dict):
        for key, item in value.items():
            yield from _paths(item, f'{prefix}.{key}' if prefix else key)
    elif isinstance(value, list) and value and isinstance(value[0], dict):
        yield from _paths(value[0], prefix + '[]')
    elif not isinstance(value, list):  # empty lists have no fields to compare
        yield prefix


def test_example_context_covers_a_real_report(service, trio_db):
    """The example used by the check must have every field the report builder produces."""
    for report in service.process({'dbpath': trio_db, 'mode': 'lpwgs'})['reports']:
        real, known = set(_paths(report['context'])), set(tpl.field_paths(report['kind']))
        assert real <= known, real - known


def test_add_list_render_choose_delete(service, single_db):
    template = add(service, 'НМИЦ 2026', 'Лаборатория НМИЦ')
    assert template['id'].startswith('u-') and template['check']['ok']

    listing = service.list_templates({})
    assert [t['id'] for t in listing['templates']][:3] == ['DZM', 'FND', '10x']
    assert listing['templates'][-1]['title'] == 'НМИЦ 2026' and listing['editable']

    processed = service.process({'dbpath': single_db})
    assert {'id': template['id'], 'title': 'НМИЦ 2026'} in processed['layouts']['default']['templates']
    assert processed['layouts']['default']['template_default'] == 'DZM'

    rendered = service.render(
        {'context': processed['reports'][0]['context'], 'kind': 'default', 'template': template['id']}
    )
    assert 'Лаборатория НМИЦ' in '\n'.join(p.text for p in Document(io.BytesIO(rendered.data)).paragraphs)

    service.update_settings({'template': template['id']})
    assert service.process({'dbpath': single_db})['layouts']['default']['template_default'] == template['id']
    assert service.template_file({'id': template['id']}).filename == 'НМИЦ 2026.docx'

    service.delete_template({'id': template['id']})
    assert template['id'] not in [t['id'] for t in service.list_templates({})['templates']]
    # the user's default pointed at it: the choice falls back to the built-in default
    assert service.process({'dbpath': single_db})['layouts']['default']['template_default'] == 'DZM'


def test_add_rejects_bad_input(service):
    with pytest.raises(service.ClientError, match='уже есть'):
        add(service, 'dzm')
    with pytest.raises(service.ClientError, match='Ошибка в полях'):
        add(service, 'Сломанный', '{{ patient.id ')
    with pytest.raises(service.ClientError, match='повреждён'):
        service.add_template({'title': 'X', 'kind': 'default', 'data': '%%%'})
    with pytest.raises(service.ClientError, match='Встроенные'):
        service.delete_template({'id': 'DZM'})


def test_unknown_fields_are_stored_with_a_warning(service):
    item = add(service, 'Черновик', '{{ patient.birth_dat }}')
    assert item['check']['ok'] is False and item['check']['unknown'][0]['field'] == 'patient.birth_dat'


def test_replace_keeps_the_id(service):
    first = add(service, 'Педиатрия', '{{ patient.x }}')
    second = add(service, 'Педиатрия', 'ok', replace=first['id'])
    assert second['id'] == first['id'] and second['check']['ok']
    assert len([t for t in service.list_templates({})['templates'] if not t['builtin']]) == 1


def test_kind_mismatch_is_rejected(service, trio_db):
    tenx = add(service, 'LPWGS лаб', kind='10x', base='10x')
    report = service.process({'dbpath': trio_db, 'mode': 'lpwgs'})['reports'][0]
    with pytest.raises(service.ClientError, match='не подходит'):
        service.render({'context': report['context'], 'kind': 'default', 'template': tenx['id']})
    with pytest.raises(service.ClientError, match='стандартного'):
        service.update_settings({'template': tenx['id']})


def test_only_admins_change_templates_in_login_mode(service):
    bob = service.Caller('bob', multiuser=True)
    admin = service.Caller('admin', multiuser=True)
    assert service.list_templates({}, bob)['editable'] is False
    assert service.get_settings({}, bob)['templates_editable'] is False
    with pytest.raises(service.ClientError, match='администратор'):
        add(service, 'X', caller=bob)
    item = add(service, 'Общий', caller=admin)
    with pytest.raises(service.ClientError, match='администратор'):
        service.delete_template({'id': item['id']}, bob)
    # everyone may use and check it
    assert service.check_template({'id': item['id']}, bob)['ok']
    service.update_settings({'template': item['id']}, bob)


def test_desktop_edition_has_no_database_features(service, single_db, monkeypatch):
    assert service.get_settings({})['edition'] == 'server'
    monkeypatch.setenv('CLINREPORT_EDITION', 'desktop')
    assert service.get_settings({})['edition'] == 'desktop'
    layouts = service.process({'dbpath': single_db})['layouts']
    assert not layouts['default']['upload'] and not layouts['default']['similar']


def test_fields_listing(service):
    fields = service.template_fields({})
    assert 'patient.id' in fields['default'] and 'p_variants[].gene' in fields['default']
    assert 'main_variants[].hgvsg' in fields['10x']


def test_conf_dir_override(service, monkeypatch, tmp_path):
    monkeypatch.setenv('CLINREPORT_CONF_DIR', str(tmp_path / 'elsewhere'))
    service.update_settings({'clinician': 'A'})
    assert (tmp_path / 'elsewhere' / 'clinreport.json').is_file()


def test_ru_status_and_refresh(monkeypatch):
    from clinreport_core import service

    monkeypatch.setattr(
        service,
        'get_ru_annotations',
        lambda: {
            'omim': {'Заболевание': {'A': 'x', 'B': 'y'}},
            'secondary': {'Disease': {'C': 'z'}, 'Inh': {'C': 'AR'}},
        },
    )
    for key, value in (('at', 0.0), ('data', None), ('warning', None)):
        monkeypatch.setitem(service._ru_cache, key, value)
    assert service.ru_status({}) == {'loaded_at': None, 'counts': {}, 'warning': None}
    status = service.ru_refresh({})
    assert status['counts'] == {'omim': 2, 'secondary': 1} and status['loaded_at'] and status['warning'] is None


def test_large_template_upload_over_aiohttp(service):
    """aiohttp's default 1 MB body limit must not block a template with pictures."""
    pytest.importorskip('aiohttp')
    from aiohttp import web
    from aiohttp.test_utils import TestClient, TestServer
    from helpers import load_route_module

    padded = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx_with('x'))) as src, zipfile.ZipFile(padded, 'w') as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr('word/media/padding.bin', os.urandom(1_500_000), compress_type=zipfile.ZIP_STORED)

    async def check():
        route = load_route_module()
        app = web.Application()  # default client_max_size = 1 MB, like the OpenCRAVAT 2.x GUI
        for method, path, handler in route.routes:
            app.router.add_route(method, '/webapps/clinreport/' + path, handler)
        async with TestClient(TestServer(app)) as client:
            response = await client.post(
                '/webapps/clinreport/template_add',
                json={'title': 'С картинкой', 'kind': 'default', 'data': b64(padded.getvalue())},
            )
            assert response.status == 200, await response.text()

    asyncio.run(check())
