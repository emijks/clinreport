#! /usr/bin/env python3
"""End-to-end smoke test of a running `oc gui` with the ClinReport webapp installed.

For every result file: open -> process (standard and LPWGS) -> edit a field ->
render every template -> check the .docx content. Uses only the HTTP API, so it
exercises route.py exactly as the review page does.

    <opencravat python> tests/e2e/smoke_http.py RESULT.sqlite [RESULT2.sqlite ...]
        [--base http://localhost:8080] [--out DIR]

Needs python-docx (installed with docxtpl). Exit code 1 on any failed check.
"""

import argparse
import io
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request

from docx import Document

EDIT_MARK = 'SMOKE-TEST EDIT'


def call(base, method, action, payload=None):
    url = f'{base}/webapps/clinreport/{action}'
    if method == 'GET':
        url += '?' + urllib.parse.urlencode(payload or {})
        request = urllib.request.Request(url)
    else:
        request = urllib.request.Request(url, json.dumps(payload).encode(), {'Content-Type': 'application/json'})
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return response.read()
    except urllib.error.HTTPError as e:
        sys.exit(f'{action}: HTTP {e.code}: {e.read().decode("utf-8", "replace")}')


def docx_text(data):
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return '\n'.join(parts)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('results', nargs='+', help='OpenCRAVAT result .sqlite files, as paths on the GUI server')
    parser.add_argument('--base', default='http://localhost:8080')
    parser.add_argument('--out', help='also save the rendered .docx files here')
    args = parser.parse_args()

    failures = 0
    for dbpath in args.results:
        if os.path.exists(dbpath):  # local file; otherwise a path on the server (e.g. inside a container)
            dbpath = os.path.abspath(dbpath)
        opened = json.loads(call(args.base, 'GET', 'open', {'dbpath': dbpath}))
        print(f'== {opened["filename"]}: samples {[s["label"] for s in opened["samples"]]}')
        modes = ['default', 'lpwgs'] if len(opened['samples']) > 1 else ['default']
        for mode in modes:
            result = json.loads(call(args.base, 'POST', 'process', {'dbpath': dbpath, 'target': 0, 'mode': mode}))
            for warning in result['warnings']:
                print(f'   warning: {warning}')
            for report in result['reports']:
                context = report['context']
                if report['kind'] == 'default':
                    context['patient']['diagnosis'] = EDIT_MARK
                templates = ['10x'] if report['kind'] == '10x' else ['DZM', 'FND']
                genes = sorted(
                    {
                        row['gene']
                        for value in context.values()
                        if isinstance(value, list)
                        for row in value
                        if isinstance(row, dict) and row.get('gene')
                    }
                )
                for template in templates:
                    data = call(
                        args.base,
                        'POST',
                        'render',
                        {
                            'context': context,
                            'kind': report['kind'],
                            'template': template,
                            'filename': report['filename'],
                        },
                    )
                    text = docx_text(data)
                    checks = {
                        'docx': data[:2] == b'PK',
                        'genes': all(g in text for g in genes),
                        'edit': report['kind'] == '10x' or EDIT_MARK in text,
                        'no None': 'None' not in text,
                    }
                    ok = all(checks.values())
                    failures += not ok
                    name = report['filename'].replace('.docx', f' [{template}].docx')
                    print(f'   [{mode}] {name:45s} {"OK  " if ok else "FAIL"} {checks}')
                    if args.out:
                        os.makedirs(args.out, exist_ok=True)
                        with open(os.path.join(args.out, name), 'wb') as fh:
                            fh.write(data)
    if failures:
        sys.exit(f'{failures} check(s) failed')
    print('all checks passed')


if __name__ == '__main__':
    main()
