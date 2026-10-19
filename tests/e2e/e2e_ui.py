#! /usr/bin/env python3
"""Browser test of the whole clinician flow against a running `oc gui`.

Viewer -> Summary-tab widget -> review page (popup) -> Обработать -> edit ->
template FND -> Сохранить как… (download) -> second tab -> LPWGS -> Настройки.
Saves a screenshot per step and the downloaded .docx, and checks the edits
reached the document.

    python tests/e2e/e2e_ui.py RESULT.sqlite --out DIR [--base http://localhost:8080] [--headed]

Needs `pip install playwright python-docx` in any Python (not necessarily
OpenCRAVAT's). Uses the installed Microsoft Edge (channel='msedge'), so no
browser download is needed; pass --channel chrome for Chrome.

The native "Save as" picker cannot be driven by automation, so the script
removes window.showSaveFilePicker and the page falls back to a download. The
picker path must be checked by hand (see docs/08-testing.md).
"""

import argparse
import os
import sys

from docx import Document
from playwright.sync_api import sync_playwright

DIAGNOSIS = 'E2E: диагноз из браузера'
CUSTOM_MARK = 'E2E: свой шаблон лаборатории'
TEMPLATES_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', 'modules', 'webapps', 'clinreport', 'templates'
)


def make_template(path, text):
    """DZM with one extra paragraph."""
    document = Document(os.path.join(TEMPLATES_DIR, 'DZM.docx'))
    document.add_paragraph(text)
    document.save(path)


DISEASE = 'E2E: заболевание из браузера'


def docx_text(path):
    document = Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return '\n'.join(parts)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('result', help='multi-sample result .sqlite (for the LPWGS step), as a path on the GUI server')
    parser.add_argument('--out', required=True, help='directory for screenshots and the .docx')
    parser.add_argument('--base', default='http://localhost:8080')
    parser.add_argument('--channel', default='msedge')
    parser.add_argument('--headed', action='store_true')
    args = parser.parse_args()
    os.makedirs(args.out, exist_ok=True)
    shot = lambda page, name, **kw: page.screenshot(path=os.path.join(args.out, name + '.png'), **kw)  # noqa: E731
    errors = []

    with sync_playwright() as p:
        browser = p.chromium.launch(channel=args.channel, headless=not args.headed)
        context = browser.new_context(viewport={'width': 1400, 'height': 900}, accept_downloads=True)
        context.add_init_script('delete window.showSaveFilePicker')  # before any page loads
        viewer = context.new_page()
        viewer.on('pageerror', lambda e: errors.append(f'viewer: {e}'))
        result = (
            os.path.abspath(args.result) if os.path.exists(args.result) else args.result
        )  # else: path on the server
        viewer.goto(f'{args.base}/result/index.html?dbpath={result}')

        # The widget renders once the viewer finished loading (can take a while on big results).
        button = viewer.get_by_text('Открыть окно заключения')
        button.first.wait_for(state='attached', timeout=90000)
        viewer.wait_for_timeout(1000)
        shot(viewer, '1_summary_widget')

        with context.expect_page() as popup:
            button.first.click()
        page = popup.value
        page.on('pageerror', lambda e: errors.append(f'app: {e}'))
        page.on('console', lambda m: m.type == 'error' and errors.append(f'app console: {m.text}'))
        page.wait_for_selector('#process-card:not([hidden])')
        shot(page, '2_review_open')

        page.click('#process-btn')
        page.wait_for_selector('#reports:not([hidden]) .panel')
        page.wait_for_timeout(500)
        shot(page, '3_processed', full_page=True)

        page.locator('.common-grid input').nth(3).fill(DIAGNOSIS)  # Предварительный диагноз
        disease = page.locator('td.col-disease textarea')
        has_disease_cell = disease.count() > 0
        if has_disease_cell:
            disease.first.fill(DISEASE)
        page.select_option('.panel-actions select', 'FND')
        with page.expect_download() as download_info:
            page.click('text=Сохранить как…')
        download = download_info.value
        target = os.path.join(args.out, download.suggested_filename)
        download.save_as(target)
        page.wait_for_timeout(500)
        shot(page, '4_saved')

        if page.locator('.tab').count() > 1:
            page.locator('.tab').nth(1).click()
            page.wait_for_timeout(300)
            shot(page, '5_second_tab', full_page=True)
            page.click('#process-lpwgs-btn')
            page.wait_for_timeout(1500)
            page.locator('.tab').nth(1).click()
            page.wait_for_timeout(300)
            shot(page, '6_lpwgs_10x', full_page=True)

        page.click('#settings-btn')
        page.wait_for_selector('#settings-dialog[open]')
        shot(page, '7_settings')
        page.click('#settings-cancel')

        # Custom templates: add one, use it for the first report; add one with a typo and see the warning.
        page.locator('.tab').first.click()
        good, bad = os.path.join(args.out, 'Свой E2E.docx'), os.path.join(args.out, 'Опечатка E2E.docx')
        make_template(good, CUSTOM_MARK)
        make_template(bad, '{{ patient.birth_dat }}')
        page.click('#templates-btn')
        page.wait_for_selector('#templates-dialog[open] .tpl-row')
        template_added = template_warned = False
        if page.locator('#template-add-form:not([hidden])').count():
            for path in (good, bad):
                page.set_input_files('#tpl-file', path)
                page.click('#tpl-add-btn')
                page.wait_for_selector(f'.tpl-name:text-is("{os.path.splitext(os.path.basename(path))[0]}")')
            template_added = True
            template_warned = page.locator('.tpl-check.warn code', has_text='patient.birth_dat').count() > 0
            shot(page, '8_templates', full_page=True)
            page.click('#templates-close')
            page.select_option('.panel-actions select', label='Свой E2E')
            with page.expect_download() as download_info:
                page.click('text=Сохранить как…')
            custom_target = os.path.join(args.out, 'custom-' + download_info.value.suggested_filename)
            download_info.value.save_as(custom_target)
            # clean up so the test can run again
            page.click('#templates-btn')
            for name in ('Свой E2E', 'Опечатка E2E'):
                row = page.locator('.tpl-row', has=page.locator(f'.tpl-name:text-is("{name}")'))
                row.get_by_role('button', name='Удалить').click()
                row = page.locator('.tpl-row', has=page.locator(f'.tpl-name:text-is("{name}")'))
                row.get_by_role('button', name='Удалить').click()
                page.wait_for_selector(f'.tpl-name:text-is("{name}")', state='detached')
        browser.close()

    text = docx_text(target)
    checks = {
        'download': os.path.getsize(target) > 0,
        'diagnosis edit in docx': DIAGNOSIS in text,
        'disease edit in docx': DISEASE in text if has_disease_cell else True,
        'custom template added': template_added,
        'unknown field warned': template_warned,
        'custom template used': template_added and CUSTOM_MARK in docx_text(custom_target),
        'no page errors': not errors,
    }
    for name, ok in checks.items():
        print(f'{"OK  " if ok else "FAIL"} {name}')
    for error in errors:
        print('   ', error)
    print(f'screenshots and {os.path.basename(target)} in {args.out}')
    if not all(checks.values()):
        sys.exit(1)


if __name__ == '__main__':
    main()
